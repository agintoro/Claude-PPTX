#!/usr/bin/env python3
"""Generate DRIVA Processing Vol. 01 — PPTX presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Color palette ─────────────────────────────────────────────────────────────
BG          = RGBColor(0x0E, 0x13, 0x08)   # very dark green-black
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE   = RGBColor(0xD8, 0xD8, 0xC8)
BULLET_TXT  = RGBColor(0xC8, 0xC8, 0xB8)
TAG_COLOR   = RGBColor(0x72, 0x7A, 0x64)
BULLET_MARK = RGBColor(0x62, 0x74, 0x4E)
FOOTER_CLR  = RGBColor(0x58, 0x5C, 0x50)
LINE_CLR    = RGBColor(0x2A, 0x30, 0x22)

# ── Slide dimensions: 10 × 10 inches (square, matches 1080×1080 source) ───────
W  = Inches(10)
H  = Inches(10)
ML = Inches(0.55)
MR = Inches(0.55)
CW = W - ML - MR

# ── Content ───────────────────────────────────────────────────────────────────
SLIDES = [
    dict(
        tag      = "— DRIVA · FERMENTATION",
        headline = "Something happens\nin the dark.",
        subtitle = "120 hours. No shortcuts. No guessing.\nOnly the process — speaking through flavor",
        bullets  = [
            "Acidity held within a precise, deliberate rhythm",
            "Pressure logged — never left to chance",
            "Temperature follows logic, not luck",
            "Flavor is not a mystery. Flavor is a decision.",
        ],
        footer   = "DRIVA FERMENTATION ARCHITECTURE",
        tagline  = "FERMENTATION ARCHITECTURE",
        page     = "1 / 4",
    ),
    dict(
        tag      = "— DRIVA · FLAVOR",
        headline = "At 1,600 to 1,800\nmetres above the sea —",
        subtitle = "The air is thin. The nights are cold.\nAnd the coffee takes its time.",
        bullets  = [
            "High altitude slows everything — and that's the point",
            "Longer cherry maturation — deeper sugars, layered complexity",
            "Cool nights preserve acidity, florals, and delicate fruit",
            "Apricot. Passion Fruit. Mandarin. Born here. Shaped by us.",
        ],
        footer   = "SIDRAP COLLECTIVE  ·  BARABABALI  ·  ATENG SUPER  ·  TONKA  ·  YELLOW CATIMOR",
        tagline  = "TRACING COFFEE TO ITS SOUL",
        page     = "2 / 4",
    ),
    dict(
        tag      = "— DRIVA · PHILOSOPHY",
        headline = "We don't sell\nbeautiful coffee.",
        subtitle = "We document\nhow that beauty is made.",
        bullets  = [
            "Every lot carries its own history",
            "Every process leaves a trail that can be read",
            "Flavor DNA — not a claim, but a record",
            "Honest coffee needs no exaggeration",
        ],
        footer   = "CLASSIC WASHED  ·  FULL WASHED  ·  THERMAL DRYER  ·  LIQUID YEAST INOCULATION",
        tagline  = "DATA-DRIVEN PROCESSING",
        page     = "3 / 4",
    ),
    dict(
        tag      = "— DRIVA · ORIGIN",
        headline = "There is a reason\nwe chose this place.",
        subtitle = "Enrekang. Barababali.\nNot because it's close — because it's right.",
        bullets  = [
            "Nights drop to 13°C here — sugars develop slowly, deeply",
            "A land and its varieties that cannot be replicated",
            "Ateng Super. Typica. Yellow Caturra — each with its own voice",
            "Terroir is not the backdrop. It is the protagonist.",
        ],
        footer   = "ENREKANG COLLECTIVE  ·  BARABABALI  ·  MARADANG VILLAGE  ·  WEST JAVA",
        tagline  = "NEXT LEVEL TERROIR",
        page     = "4 / 4",
    ),
]


def _run(para, text, name, size, color, bold=False, italic=False):
    r = para.add_run()
    r.text = text
    r.font.name = name
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.italic = italic
    return r


def _tb(slide, left, top, width, height, wrap=True):
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text_frame.word_wrap = wrap
    return shape.text_frame


def _rect(slide, left, top, width, height, color):
    shp = slide.shapes.add_shape(1, left, top, width, height)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.width = Emu(0)
    return shp


def make_slide(prs, data):
    sl = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    # Background
    bg = sl.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG

    # ── 1. Category tag ───────────────────────────────────────────────────────
    tf = _tb(sl, ML, Inches(0.40), CW, Inches(0.25))
    _run(tf.paragraphs[0], data["tag"], "Arial", 7.5, TAG_COLOR)

    # ── 2. Headline ───────────────────────────────────────────────────────────
    tf = _tb(sl, ML, Inches(0.82), CW, Inches(3.80))
    for i, line in enumerate(data["headline"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(80)
        _run(p, line, "Times New Roman", 72, WHITE, bold=True)

    # ── 3. Subtitle ───────────────────────────────────────────────────────────
    tf = _tb(sl, ML, Inches(4.76), CW, Inches(0.88))
    for i, line in enumerate(data["subtitle"].split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = Pt(19)
        _run(p, line, "Arial", 12, OFF_WHITE)

    # ── 4. Bullet points ──────────────────────────────────────────────────────
    y = Inches(5.76)
    for bullet in data["bullets"]:
        tf = _tb(sl, ML, y, CW, Inches(0.44))
        p = tf.paragraphs[0]
        _run(p, "·  ", "Arial", 11, BULLET_MARK)
        _run(p, bullet, "Arial", 11, BULLET_TXT)
        y += Inches(0.50)

    # ── 5. Separator line ─────────────────────────────────────────────────────
    _rect(sl, ML, Inches(7.90), CW, Emu(12700), LINE_CLR)  # 1 pt tall

    # ── 6. Footer tags ────────────────────────────────────────────────────────
    tf = _tb(sl, ML, Inches(8.04), CW, Inches(0.28))
    _run(tf.paragraphs[0], data["footer"], "Arial", 6.5, FOOTER_CLR)

    # ── 7. Brand bar ──────────────────────────────────────────────────────────
    brand_y = Inches(8.74)

    tf = _tb(sl, ML, brand_y, Inches(1.3), Inches(0.36))
    _run(tf.paragraphs[0], "DRIVA", "Arial", 12, WHITE, bold=True)

    tf = _tb(sl, ML + Inches(1.3), brand_y, Inches(6.3), Inches(0.36))
    _run(tf.paragraphs[0], data["tagline"], "Arial", 7, FOOTER_CLR)

    tf = _tb(sl, W - MR - Inches(0.80), brand_y, Inches(0.80), Inches(0.36))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    _run(p, data["page"], "Arial", 7.5, FOOTER_CLR)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    for data in SLIDES:
        make_slide(prs, data)

    out = "DRIVA Processing Vol. 01.pptx"
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    main()
