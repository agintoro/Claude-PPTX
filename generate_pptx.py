#!/usr/bin/env python3
"""Generate Partners Thank You.pptx for Driva Coffee Co."""

import io, os
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.oxml.ns import qn
from lxml import etree

def px(n): return Emu(int(round(n * 9525)))

WARM  = RGBColor(0xd8, 0xd0, 0xbf)
INK   = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x6b, 0x66, 0x5d)
CARD  = RGBColor(0xff, 0xff, 0xff)
LINE  = RGBColor(0xe7, 0xdf, 0xd0)
PAPER = RGBColor(0xf7, 0xf3, 0xea)

SLIDE_W, SLIDE_H = 1920, 1080
PAD_X, PAD_Y = 90, 70
CONTENT_W = SLIDE_W - 2 * PAD_X  # 1740
DISC = 160     # logo circle diameter (px)
NAME_H = 32    # height for name + city text block
GAP_Y = 36     # vertical gap between rows

BASE_DIR  = "/home/user/Claude-PPTX/pptx-partners/project"
LOGOS_DIR = os.path.join(BASE_DIR, "logos")

PARTNERS = [
    {"file": "rumah-sangrai-bloom.png", "name": "Rumah Sangrai\nBloom", "city": "Bandung",             "scale": 1.0},
    {"file": "forth.png",               "name": "Forth Coffee Co.",     "city": "Jakarta",              "scale": 1.0},
    {"file": "hana.png",                "name": "Hana Roastery",        "city": "Surabaya",             "scale": 1.15},
    {"file": "n.png",                   "name": "N. Roasters",          "city": "Yogyakarta",           "scale": 1.0},
    {"file": "contrast.png",            "name": "Contrast",             "city": "Roastery — Bali",      "scale": 1.35},
    {"file": "kozi.png",                "name": "Kozi Company",         "city": "Jakarta",              "scale": 1.0},
    {"file": "koeslan.png",             "name": "Koeslan's Coffee",     "city": "Semarang",             "scale": 1.55},
    {"file": "instinct.png",            "name": "Instinct Roastery",    "city": "Bandung",              "scale": 1.55},
    {"file": "fugol.png",               "name": "Fugol",                "city": "Roasters — Bali",      "scale": 1.0},
    {"file": "froast.png",              "name": "F. Roast Assembly",    "city": "Surabaya",             "scale": 1.0},
    {"file": "steambrew.png",           "name": "Steam & Brew",         "city": "Specialty — Jakarta",  "scale": 1.15},
    {"file": "kofind.png",              "name": "Kofind Coffee Co.",    "city": "Bandung",              "scale": 1.35},
    {"file": "peopletemple.png",        "name": "People Temple",        "city": "Jakarta",              "scale": 1.0},
    {"file": "kopi-oemah-ladi.png",     "name": "Kopi Oemah l’adi","city": "Solo",                "scale": 1.35},
]


def make_circle_logo(img_path, disc=160, css_scale=1.0):
    """Return a PIL RGB image of the logo clipped to a circle with white fill."""
    HR = 4
    cs = disc * HR
    img = Image.open(img_path).convert("RGBA")

    # object-fit: cover — scale so both dimensions >= cs, then center-crop
    cov = max(cs / img.width, cs / img.height)
    w2, h2 = int(img.width * cov), int(img.height * cov)
    img = img.resize((w2, h2), Image.LANCZOS)
    l, t = (w2 - cs) // 2, (h2 - cs) // 2
    img = img.crop((l, t, l + cs, t + cs))

    # CSS transform: scale(css_scale) — zoom-in from center
    if css_scale != 1.0:
        zs = int(cs * css_scale)
        img = img.resize((zs, zs), Image.LANCZOS)
        lz, tz = (zs - cs) // 2, (zs - cs) // 2
        img = img.crop((lz, tz, lz + cs, tz + cs))

    # Flatten onto white
    white = Image.new("RGB", (cs, cs), (255, 255, 255))
    if img.mode == "RGBA":
        white.paste(img.convert("RGB"), mask=img.split()[3])
    else:
        white.paste(img.convert("RGB"))

    # Circular mask → composite on warm background colour
    mask = Image.new("L", (cs, cs), 0)
    ImageDraw.Draw(mask).ellipse([0, 0, cs - 1, cs - 1], fill=255)
    bg = Image.new("RGB", (cs, cs), (0xd8, 0xd0, 0xbf))
    bg.paste(white, mask=mask)

    # Downsample at 2× for crisp PNG in the PPTX
    return bg.resize((disc * 2, disc * 2), Image.LANCZOS)


def to_stream(img):
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def set_no_line(shape):
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, fill_rgb):
    s = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, px(x), px(y), px(w), px(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill_rgb
    set_no_line(s)
    return s


def add_tb(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, wrap=False):
    """Add a textbox. runs = list of (text, pt_size, bold, italic, color)."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    # kill default margin
    from pptx.util import Inches
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    for text, size, bold, italic, color in runs:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = "Arial"
        r.font.color.rgb = color
    return tb


def add_multiline_tb(slide, x, y, w, h, paras, wrap=False):
    """paras = list of (align, runs_list)."""
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    for i, (align, runs) in enumerate(paras):
        p = tf.paragraphs[i] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for text, size, bold, italic, color in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = "Arial"
            r.font.color.rgb = color
    return tb


# ── Build presentation ────────────────────────────────────────────────────────

prs = Presentation()
prs.slide_width  = px(SLIDE_W)
prs.slide_height = px(SLIDE_H)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout

# Background
add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, WARM)

# ── Eyebrow ───────────────────────────────────────────────────────────────────
add_tb(slide, PAD_X, PAD_Y + 22, CONTENT_W - 80, 22,
       [("DRIVA COFFEE CO.   ·   INDRAGIRI / PATENGAN / CITAMIANG   ·   2026 SEASON",
         8.5, True, False, MUTED)])

# ── D mark (rounded rectangle, ink fill, white "D") ──────────────────────────
MARK_X = SLIDE_W - PAD_X - 64
mark = slide.shapes.add_shape(
    MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    px(MARK_X), px(PAD_Y), px(64), px(64)
)
mark.fill.solid()
mark.fill.fore_color.rgb = INK
set_no_line(mark)
# Roundness: 18px radius on 64px → ratio ≈ 0.28
mark.adjustments[0] = 0.28
tf = mark.text_frame
tf.word_wrap = False
tf.margin_left = tf.margin_right = Emu(0)
tf.margin_top = Emu(px(12).emu if hasattr(px(12), 'emu') else int(12*9525))
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "D"
r.font.size = Pt(27)
r.font.bold = True
r.font.name = "Arial"
r.font.color.rgb = PAPER

# ── Headline (y = PAD_Y + mark_height + 36 = 170) ────────────────────────────
HEAD_Y = PAD_Y + 64 + 36   # 170
DISP_H = 200                # display text block height
DISP_W = int(CONTENT_W * 0.62)  # ≈ 1078

add_multiline_tb(
    slide, PAD_X, HEAD_Y, DISP_W, DISP_H,
    [
        (PP_ALIGN.LEFT, [("Terima kasih,", 75, True, False, INK)]),
        (PP_ALIGN.LEFT, [
            ("to our ", 75, False, True, MUTED),
            ("partners.", 75, True, False, INK),
        ]),
    ]
)

# ── Subhead ───────────────────────────────────────────────────────────────────
SUB_X = PAD_X + DISP_W + 32
SUB_W = SLIDE_W - PAD_X - SUB_X
sub_text = (
    "To the fourteen roasters who put our lots\n"
    "on the bar this season — sorted, fermented,\n"
    "dried and rested by hand.\n"
    "Watch the night RH. Keep the airflow steady."
)
add_tb(slide, SUB_X, HEAD_Y + 80, SUB_W, 140,
       [(sub_text, 12, False, False, MUTED)],
       align=PP_ALIGN.RIGHT, wrap=True)

# ── Divider rule ──────────────────────────────────────────────────────────────
RULE_Y = HEAD_Y + 186 + 40   # ≈ 396

add_tb(slide, PAD_X, RULE_Y - 4, 160, 16,
       [("PARTNER ROASTERS", 7, True, False, MUTED)])

add_rect(slide, PAD_X + 168, RULE_Y + 5, CONTENT_W - 248, 1, LINE)

add_tb(slide, SLIDE_W - PAD_X - 72, RULE_Y - 4, 72, 16,
       [("14 / 14", 7, True, False, INK)], align=PP_ALIGN.RIGHT)

# ── Partner logo grid ─────────────────────────────────────────────────────────
GRID_Y = RULE_Y + 12 + 48   # ≈ 456

COLS   = 7
COL_W  = CONTENT_W // COLS           # 248
ROW_H  = DISC + 18 + NAME_H          # 210

for i, pd in enumerate(PARTNERS):
    col = i % COLS
    row = i // COLS

    col_cx = PAD_X + col * COL_W + COL_W // 2
    disc_x = col_cx - DISC // 2
    disc_y = GRID_Y + row * (ROW_H + GAP_Y)

    logo_img = make_circle_logo(
        os.path.join(LOGOS_DIR, pd["file"]),
        disc=DISC,
        css_scale=pd["scale"],
    )
    slide.shapes.add_picture(to_stream(logo_img), px(disc_x), px(disc_y), px(DISC), px(DISC))

    name_y = disc_y + DISC + 14
    name_x = disc_x - 24
    name_w = DISC + 48

    add_multiline_tb(
        slide, name_x, name_y, name_w, NAME_H,
        [
            (PP_ALIGN.CENTER, [(pd["name"], 7, True, False, INK)]),
            (PP_ALIGN.CENTER, [(pd["city"],  6, True, False, MUTED)]),
        ]
    )

# ── Footer ────────────────────────────────────────────────────────────────────
FOOT_Y = SLIDE_H - PAD_Y - 34

add_rect(slide, PAD_X, FOOT_Y, CONTENT_W, 1, LINE)

FT_Y = FOOT_Y + 18
add_tb(slide, PAD_X, FT_Y, 500, 14,
       [("DRIVA PROCESSING OS · SEASON 2026", 7, True, False, MUTED)])

add_tb(slide, PAD_X + 500, FT_Y, CONTENT_W - 900, 14,
       [("DRV-IND-2605 · DRV-PAT-2605 · DRV-CIT-2605", 8, True, False, INK)],
       align=PP_ALIGN.CENTER)

add_tb(slide, SLIDE_W - PAD_X - 200, FT_Y, 200, 14,
       [("WITH THANKS.", 7, True, False, MUTED)],
       align=PP_ALIGN.RIGHT)

# ── Save ──────────────────────────────────────────────────────────────────────
OUT = "/home/user/Claude-PPTX/Partners Thank You.pptx"
prs.save(OUT)
print(f"Saved → {OUT}")
